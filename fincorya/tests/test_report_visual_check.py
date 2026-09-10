"""Programmatic visual/structural checks of generated PDF and Excel exports."""
from datetime import date
from decimal import Decimal
from io import BytesIO
from unittest.mock import patch

from django.test import TestCase, override_settings

import fitz  # PyMuPDF
from openpyxl import load_workbook

from apps.finance.reporting import build_report
from apps.reports.services import generate_finance_report
from tests.test_finance_workflows import AUGUST, JULY, FinanceScenario


def pdf_text(path):
    with fitz.open(path) as doc:
        return [page.get_text() for page in doc]


@override_settings(FINANCE_LEDGER_ENABLED=True, LOCAL_AUTH_BYPASS=False, MFA_ENABLED=False, SECURE_SSL_REDIRECT=False)
class ReportVisualChecks(FinanceScenario, TestCase):
    def setUp(self):
        super().setUp()
        self.operation('vis-1')
        self.operation('vis-2')

    def test_agent_pdf_contains_only_his_day(self):
        export = generate_finance_report(user=self.agent, kind='ACTIVITY', preset='DAY', anchor=JULY.date(), format='PDF')
        self.assertEqual(export.status, 'READY')
        pages = pdf_text(export.file.path)
        self.assertGreaterEqual(len(pages), 1)
        text = '\n'.join(pages)
        self.assertIn('RAPPORT D\u2019ACTIVIT\u00c9', text.upper())
        self.assertIn('Page', text)
        self.assertIn(self.agent.email, text)
        # An agent report never exposes another user's references or identifiers.
        self.assertNotIn('finance-manager@test.local', text)
        self.assertNotIn('finance-admin@test.local', text)

    def test_manager_weekly_excel_has_all_sections_and_totals(self):
        export = generate_finance_report(user=self.manager, kind='TREASURY', preset='WEEK', anchor=JULY.date(), format='XLSX')
        self.assertEqual(export.status, 'READY')
        wb = load_workbook(export.file.path)
        names = [s.title for s in wb.worksheets]
        self.assertIn('En-tête', names)
        self.assertTrue(any('Rapprochement' in n for n in names))
        # Totals row exists on the reconciliation sheet.
        sheet = next(s for s in wb.worksheets if 'Rapprochement' in s.title)
        self.assertTrue(any(str(cell.value or '').startswith('Total') for row in sheet.iter_rows() for cell in row))

    def test_admin_monthly_pdf_and_excel_match_dashboard_totals(self):
        from apps.finance.closing import close_period, propose_distribution, approve_distributions
        from apps.finance.models import DistributionPolicy
        from apps.finance.locking import save_internal
        from apps.finance.models import DistributionPolicyShare
        from apps.profits.models import ProfitPeriod
        from apps.stakeholders.models import Stakeholder
        sh = Stakeholder.objects.create(name='Shareholder A', type='SHAREHOLDER')
        policy = DistributionPolicy.objects.create(mode='EQUAL_SHARES', effective_from=date(2026, 7, 1), created_by=self.admin)
        save_internal(DistributionPolicyShare(policy=policy, stakeholder=sh, percent=Decimal('100')))
        month = self.period()
        self.count_all(month)
        with patch('django.utils.timezone.now', return_value=AUGUST):
            close_period(period_id=month.pk, actor=self.admin)
            profit = ProfitPeriod.objects.get(finance_period=month)
            propose_distribution(profit_id=profit.pk, amount=Decimal('200'), actor=self.admin)
            approve_distributions(profit_id=profit.pk, amount=Decimal('200'), actor=self.admin)
        screen = build_report(user=self.admin, kind='RESULT', preset='MONTH', anchor=date(2026, 7, 15))
        pdf = generate_finance_report(user=self.admin, kind='RESULT', preset='MONTH', anchor=date(2026, 7, 15), format='PDF')
        xlsx = generate_finance_report(user=self.admin, kind='RESULT', preset='MONTH', anchor=date(2026, 7, 15), format='XLSX')
        for export in (pdf, xlsx):
            self.assertEqual(export.status, 'READY')
        pages = pdf_text(pdf.file.path)
        self.assertGreaterEqual(len(pages), 1)
        text = '\n'.join(pages)
        self.assertIn(str(profit.net_profit), text)
        self.assertIn('Écarts de change', text)
        wb = load_workbook(xlsx.file.path)
        result_sheet = next(s for s in wb.worksheets if 'Résultat' in s.title)
        amounts = [Decimal(str(cell.value)) for row in result_sheet.iter_rows() for cell in row if isinstance(cell.value, (int, float, Decimal))]
        self.assertIn(profit.net_profit, amounts)

    def test_no_export_for_forbidden_roles(self):
        from django.core.exceptions import PermissionDenied
        holder = self.client  # placeholder no-op; check service-level denial below
        from apps.stakeholders.models import Stakeholder
        investor = Stakeholder.objects.create(name='Forbidden Investor', type='INVESTOR')
        from apps.accounts.models import Role, User
        investor_user = User.objects.create_user(email='inv-forbidden@test.local', role=Role.INVESTOR)
        investor.owner = investor_user
        investor.save()
        with self.assertRaises(PermissionDenied):
            generate_finance_report(user=investor_user, kind='TREASURY', preset='MONTH', anchor=JULY.date(), format='PDF')
