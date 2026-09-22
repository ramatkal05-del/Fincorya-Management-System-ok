import csv
import calendar
from datetime import date
from io import BytesIO, StringIO
from decimal import Decimal
from pathlib import Path

from django.core.files.base import ContentFile
from django.conf import settings
from django.db.models import Prefetch, Sum
from django.utils import timezone
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER, TA_JUSTIFY, TA_LEFT, TA_RIGHT
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.graphics.charts.barcharts import VerticalBarChart
from reportlab.graphics.charts.legends import Legend
from reportlab.graphics.shapes import Drawing, String
from reportlab.pdfgen import canvas
from reportlab.platypus import CondPageBreak, HRFlowable, Image, KeepTogether, PageBreak, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

from apps.accounts.models import Role
from apps.audit.services import record
from apps.contracts.models import Contract
from apps.operations.models import Operation, OperationStatus
from apps.expenses.models import Expense
from apps.finance.models import FundContribution
from apps.profits.models import Distribution, ProfitPeriod
from apps.stakeholders.models import Investment, PartnerOperation, PaymentFrequency, Stakeholder
from config.business_time import business_day_bounds
from .models import ReportExport


def _safe_row(values):
    from apps.finance.reporting import safe_cell
    return [safe_cell(value) for value in values]


def _style_workbook(workbook):
    green = "006B4F"
    for sheet in workbook.worksheets:
        sheet.sheet_view.showGridLines = False
        sheet.freeze_panes = "A2"
        if sheet.max_row:
            for cell in sheet[1]:
                cell.fill = PatternFill("solid", fgColor=green)
                cell.font = Font(color="FFFFFF", bold=True)
                cell.alignment = Alignment(vertical="center", wrap_text=True)
            sheet.row_dimensions[1].height = 28
        for column in sheet.columns:
            letter = column[0].column_letter
            max_length = max((len(str(cell.value or "")) for cell in column), default=8)
            sheet.column_dimensions[letter].width = min(max(max_length + 2, 12), 34)
            for cell in column[1:]:
                cell.alignment = Alignment(vertical="top", wrap_text=True)
        sheet.auto_filter.ref = sheet.dimensions
    workbook.properties.creator = "FINCORYA Group"
    workbook.properties.title = "Rapport financier FINCORYA"
    workbook.properties.subject = "Rapport confidentiel"


def operation_report_snapshot(*, user, start_date, end_date):
    from django.core.exceptions import PermissionDenied
    if not user.is_active or user.role not in {Role.ADMIN, Role.AGENT}:
        raise PermissionDenied("Ce rapport est réservé à l’administrateur et à l’agent concerné.")
    if user.role == Role.AGENT and (start_date != end_date or start_date > timezone.localdate()):
        raise PermissionDenied("Un agent télécharge uniquement son rapport quotidien, passé ou en cours.")
    period_start, period_end = business_day_bounds(start_date, end_date)
    operations = Operation.objects.select_related("currency", "agent").filter(
        created_at__gte=period_start,
        created_at__lt=period_end,
    )
    if user.role == Role.AGENT:
        operations = operations.filter(agent=user)
    elif user.role != Role.ADMIN:
        operations = operations.none()
    rows, by_currency = [], {}
    for operation in operations.order_by("created_at", "id"):
        rows.append({
            "reference": str(operation.reference), "date": timezone.localtime(operation.created_at).strftime("%d/%m/%Y %H:%M"),
            "type": operation.get_type_display(), "status": operation.get_status_display(),
            "amount": str(operation.amount), "currency": operation.currency.code,
            "fee": str(operation.fee), "agent": operation.agent.get_full_name() or operation.agent.email,
            "agent_email": operation.agent.email,
            "service": operation.get_service_display(), "customer_identifier": operation.customer_identifier,
            "customer_name": operation.customer_name,
        })
        # A cancelled operation's amount/fee were fully reversed: only
        # COMPLETED operations count towards the totals shown to the user.
        if operation.status != OperationStatus.COMPLETED:
            continue
        totals = by_currency.setdefault(operation.currency.code, {"amount": Decimal("0.00"), "fees": Decimal("0.00")})
        totals["amount"] += operation.amount
        totals["fees"] += operation.fee
    serialized_totals = {code: {"amount": str(values["amount"]), "fees": str(values["fees"])} for code, values in sorted(by_currency.items())}
    return {"start_date": str(start_date), "end_date": str(end_date), "rows": rows, "totals": {"count": len(rows), "by_currency": serialized_totals}}


def _csv_bytes(snapshot):
    stream = StringIO(newline="")
    writer = csv.writer(stream)
    writer.writerow(["reference", "date", "type", "service", "client", "identifiant", "status", "amount", "currency", "commission", "agent"])
    for row in snapshot["rows"]:
        writer.writerow(_safe_row([row[key] for key in ("reference", "date", "type", "service", "customer_name", "customer_identifier", "status", "amount", "currency", "fee", "agent")]))
    for currency, totals in snapshot["totals"]["by_currency"].items():
        writer.writerow(["TOTAL", "", "", "", "", "", str(snapshot["totals"]["count"]) + " opérations", totals["amount"], currency, totals["fees"], ""])
    return ("\ufeff" + stream.getvalue()).encode("utf-8")


def _xlsx_bytes(snapshot):
    workbook = Workbook()
    detail = workbook.active
    detail.title = "Détail"
    detail.append(["Référence", "Date", "Type", "Service", "Client", "Identifiant", "Statut", "Montant", "Devise", "Commission", "Agent"])
    for row in snapshot["rows"]:
        detail.append(_safe_row([row["reference"], row["date"], row["type"], row["service"], row["customer_name"], row["customer_identifier"], row["status"], Decimal(row["amount"]), row["currency"], Decimal(row["fee"]), row["agent"]]))
    summary = workbook.create_sheet("Synthèse")
    summary.append(["Indicateur", "Valeur"])
    summary.append(["Nombre d'opérations", snapshot["totals"]["count"]])
    for currency, totals in snapshot["totals"]["by_currency"].items():
        summary.append([f"Montant total {currency}", Decimal(totals["amount"])])
        summary.append([f"Frais totaux {currency}", Decimal(totals["fees"])])
    _style_workbook(workbook)
    stream = BytesIO()
    workbook.save(stream)
    return stream.getvalue()


def _report_fonts():
    candidates = []
    if settings.REPORT_FONT_DIR:
        font_dir = Path(settings.REPORT_FONT_DIR)
        candidates.append((font_dir / "arialn.ttf", font_dir / "arialnb.ttf"))
    candidates.extend([
        (settings.BASE_DIR / "static_src" / "fonts" / "arialn.ttf", settings.BASE_DIR / "static_src" / "fonts" / "arialnb.ttf"),
        (Path("C:/Windows/Fonts/arialn.ttf"), Path("C:/Windows/Fonts/arialnb.ttf")),
    ])
    for regular_path, bold_path in candidates:
        if regular_path.exists() and bold_path.exists():
            if "ArialNarrow" not in pdfmetrics.getRegisteredFontNames():
                pdfmetrics.registerFont(TTFont("ArialNarrow", str(regular_path)))
                pdfmetrics.registerFont(TTFont("ArialNarrow-Bold", str(bold_path)))
            return "ArialNarrow", "ArialNarrow-Bold"
    return "Helvetica", "Helvetica-Bold"


def _summary_chart(chart_data, usable_width):
    if not chart_data:
        return None
    rows = list(chart_data.items())
    drawing = Drawing(usable_width, 52 * mm)
    chart = VerticalBarChart()
    chart.x, chart.y, chart.width, chart.height = 13 * mm, 10 * mm, usable_width - 62 * mm, 34 * mm
    keys = ("operations", "commissions", "expenses", "distributions")
    chart.data = [[float(values.get(key, 0)) for _, values in rows] for key in keys]
    chart.categoryAxis.categoryNames = [code for code, _ in rows]
    chart.categoryAxis.labels.fontSize = 8
    chart.valueAxis.labels.fontSize = 7
    for bar, color in zip(chart.bars, ("#006B4F", "#C9A227", "#9AA7A2", "#013B36")):
        bar.fillColor = colors.HexColor(color)
    drawing.add(chart)
    legend = Legend()
    legend.x, legend.y = usable_width - 45 * mm, 38 * mm
    legend.fontSize = 8
    legend.alignment = "right"
    legend.colorNamePairs = [
        (colors.HexColor("#006B4F"), "Operations"),
        (colors.HexColor("#C9A227"), "Commissions"),
        (colors.HexColor("#9AA7A2"), "Depenses"),
        (colors.HexColor("#013B36"), "Distributions"),
    ]
    drawing.add(legend)
    drawing.add(String(usable_width / 2, 1 * mm, "Operations, commissions, depenses et distributions par devise", textAnchor="middle", fontSize=8, fillColor=colors.HexColor("#596763")))
    return drawing


def _transparent_report_logo():
    logo_path = settings.BASE_DIR / "static_src" / "img" / "fincorya-report-logo.png"
    return logo_path if logo_path.exists() else None


def _build_branded_pdf(*, title, subtitle, sections, downloaded_by, landscape_mode=True, chart_data=None):
    stream = BytesIO()
    page_size = landscape(A4) if landscape_mode else A4
    document = SimpleDocTemplate(
        stream, pagesize=page_size, leftMargin=14 * mm, rightMargin=14 * mm,
        topMargin=27 * mm, bottomMargin=18 * mm, title=title, author="FINCORYA Group",
    )
    regular_font, bold_font = _report_fonts()
    styles = getSampleStyleSheet()
    styles.add(ParagraphStyle(name="ReportTitle", parent=styles["Title"], fontName=bold_font, fontSize=22, leading=25, alignment=TA_CENTER, textColor=colors.HexColor("#013B36"), spaceAfter=3 * mm))
    styles.add(ParagraphStyle(name="ReportMeta", parent=styles["Normal"], fontName=regular_font, fontSize=14, leading=17, alignment=TA_JUSTIFY, textColor=colors.HexColor("#596763"), spaceAfter=6 * mm))
    styles.add(ParagraphStyle(name="SectionTitle", parent=styles["Heading2"], fontName=bold_font, fontSize=14, leading=17, textColor=colors.HexColor("#006B4F"), spaceBefore=5 * mm, spaceAfter=2.5 * mm))
    styles.add(ParagraphStyle(name="Cell", parent=styles["Normal"], fontName=regular_font, fontSize=9.5, leading=11, alignment=TA_JUSTIFY, textColor=colors.HexColor("#18201F")))
    styles.add(ParagraphStyle(name="CellCompact", parent=styles["Cell"], fontSize=7.5, leading=8.5))
    styles.add(ParagraphStyle(name="CellHeader", parent=styles["Cell"], fontName=bold_font, textColor=colors.white, alignment=TA_CENTER))
    styles.add(ParagraphStyle(name="CellHeaderCompact", parent=styles["CellCompact"], fontName=bold_font, textColor=colors.white, alignment=TA_CENTER))
    styles.add(ParagraphStyle(name="CellMoney", parent=styles["Cell"], alignment=TA_RIGHT))
    styles.add(ParagraphStyle(name="CellMoneyCompact", parent=styles["CellCompact"], alignment=TA_RIGHT))
    story = []
    story.extend([Paragraph(title, styles["ReportTitle"]), Paragraph(subtitle, styles["ReportMeta"])])
    usable_width = page_size[0] - document.leftMargin - document.rightMargin
    chart = _summary_chart(chart_data or {}, usable_width)
    if chart:
        story.extend([chart, Spacer(1, 2 * mm)])
    for section_number, section in enumerate(sections, start=1):
        section_title, data, *extra = section
        numeric = set(extra[0]) if extra else set()
        story.append(CondPageBreak(34 * mm))
        story.append(Paragraph(f"{section_number}. {section_title}", styles["SectionTitle"]))
        if len(data) <= 1:
            story.append(Paragraph("Aucune donnee pour cette section.", styles["ReportMeta"]))
            continue
        rendered = []
        compact = len(data[0]) > 10
        for row_index, row in enumerate(data):
            if compact:
                style, money = (styles["CellHeaderCompact"], styles["CellHeaderCompact"]) if row_index == 0 else (styles["CellCompact"], styles["CellMoneyCompact"])
            else:
                style, money = (styles["CellHeader"], styles["CellHeader"]) if row_index == 0 else (styles["Cell"], styles["CellMoney"])
            rendered.append([Paragraph(str(value).replace("&", "&amp;").replace("<", "&lt;"), money if index in numeric else style) for index, value in enumerate(row)])
        widths = [usable_width / len(rendered[0])] * len(rendered[0])
        table = Table(rendered, colWidths=widths, repeatRows=1, hAlign="LEFT")
        table.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#006B4F")),
            ("GRID", (0, 0), (-1, -1), .35, colors.HexColor("#DCE3DF")),
            ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ("LEFTPADDING", (0, 0), (-1, -1), 4), ("RIGHTPADDING", (0, 0), (-1, -1), 4),
            ("TOPPADDING", (0, 0), (-1, -1), 4), ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
            ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F7F8F5")]),
        ]))
        story.extend([table, Spacer(1, 3 * mm)])

    downloader = downloaded_by.get_full_name() or downloaded_by.email

    def header_footer(pdf_canvas, doc):
        pdf_canvas.saveState()
        pdf_canvas.setFont(bold_font, 9)
        pdf_canvas.setFillColor(colors.HexColor("#013B36"))
        page_logo = _transparent_report_logo()
        if page_logo:
            pdf_canvas.drawImage(
                str(page_logo), doc.leftMargin, page_size[1] - 23 * mm,
                width=30 * mm, height=20 * mm, preserveAspectRatio=True, mask="auto",
            )
        else:
            pdf_canvas.drawString(doc.leftMargin, page_size[1] - 12 * mm, "FINCORYA GROUP")
        pdf_canvas.drawRightString(page_size[0] - doc.rightMargin, page_size[1] - 12 * mm, title.upper())
        pdf_canvas.setStrokeColor(colors.HexColor("#C9A227"))
        pdf_canvas.line(doc.leftMargin, page_size[1] - 25 * mm, page_size[0] - doc.rightMargin, page_size[1] - 25 * mm)
        pdf_canvas.setStrokeColor(colors.HexColor("#DCE3DF")); pdf_canvas.line(doc.leftMargin, 10 * mm, page_size[0] - doc.rightMargin, 10 * mm)
        pdf_canvas.setFont(regular_font, 8); pdf_canvas.setFillColor(colors.HexColor("#596763"))
        pdf_canvas.drawString(doc.leftMargin, 6.5 * mm, "FINCORYA Group - Document confidentiel")
        pdf_canvas.drawCentredString(page_size[0] / 2, 6.5 * mm, f"Telecharge par : {downloader}")
        pdf_canvas.drawRightString(page_size[0] - doc.rightMargin, 6.5 * mm, f"Page {doc.page}")
        pdf_canvas.restoreState()

    document.build(story, onFirstPage=header_footer, onLaterPages=header_footer)
    return stream.getvalue()


def _pdf_bytes(snapshot, user):
    rows = [["Date / heure", "Reference", "Type", "Service", "Client", "Identifiant", "Montant", "Commission", "Agent"]]
    for row in snapshot["rows"]:
        rows.append([row["date"], row["reference"].upper(), row["type"], row["service"], row["customer_name"] or "-", row["customer_identifier"] or "-", f"{row['amount']} {row['currency']}", f"{row['fee']} {row['currency']}", f"{row['agent']}\n{row['agent_email']}"])
    totals = [[code, values["amount"], values["fees"]] for code, values in snapshot["totals"]["by_currency"].items()]
    chart_data = {
        code: {"operations": values["amount"], "commissions": values["fees"], "expenses": 0, "distributions": 0}
        for code, values in snapshot["totals"]["by_currency"].items()
    }
    return _build_branded_pdf(
        title="Rapport des operations", subtitle=f"Periode du {snapshot['start_date']} au {snapshot['end_date']}",
        sections=[("Synthese", [["Devise", "Volume", "Commissions"], *totals]), ("Detail des operations", rows)], downloaded_by=user, landscape_mode=True, chart_data=chart_data if len(chart_data) > 1 else None,
    )


def generate_operation_report(*, user, start_date, end_date, format):
    format = format.upper()
    if format not in {"CSV", "XLSX", "PDF"}:
        raise ValueError("Format de rapport non pris en charge.")
    snapshot = operation_report_snapshot(user=user, start_date=start_date, end_date=end_date)
    export = ReportExport.objects.create(requested_by=user, kind="OPERATIONS", format=format, filters={"start_date": str(start_date), "end_date": str(end_date)}, snapshot=snapshot)
    try:
        content = _pdf_bytes(snapshot, user) if format == "PDF" else {"CSV": _csv_bytes, "XLSX": _xlsx_bytes}[format](snapshot)
        export.file.save(f"operations-{start_date}-{end_date}.{format.lower()}", ContentFile(content), save=False)
        export.status, export.completed_at = "READY", timezone.now()
        export.save(update_fields=["file", "status", "completed_at"])
    except Exception:
        if export.file:
            export.file.delete(save=False)
        export.status = "FAILED"
        export.completed_at = timezone.now()
        export.save(update_fields=["status", "completed_at"])
        raise
    record(actor=user, action="REPORT_GENERATE", instance=export, after={"format": format, "totals": snapshot["totals"]})
    return export


def monthly_financial_snapshot(*, user, year, month, agent=None, stakeholder=None, stakeholder_type=""):
    if user.role != Role.ADMIN:
        raise PermissionError("Seul l'administrateur peut générer le rapport mensuel consolidé.")
    start = date(int(year), int(month), 1)
    end = date(int(year), int(month), calendar.monthrange(int(year), int(month))[1])
    start_at, end_at = business_day_bounds(start, end)
    operations = Operation.objects.select_related("agent", "currency").filter(created_at__gte=start_at, created_at__lt=end_at)
    if agent:
        operations = operations.filter(agent=agent)
    operation_rows = [{
        "reference": str(op.reference), "date": timezone.localtime(op.created_at).strftime("%d/%m/%Y %H:%M"), "agent": op.agent.get_full_name() or op.agent.email,
        "service": op.get_service_display(), "type": op.get_type_display(), "client": op.customer_name,
        "identifier": op.customer_identifier, "amount": str(op.amount), "currency": op.currency.code,
        "commission": str(op.fee), "status": op.get_status_display(),
    } for op in operations.order_by("created_at", "id")]
    expenses = Expense.objects.select_related("currency").filter(incurred_on__range=(start, end), status="APPROVED")
    expense_rows = [{"date": str(row.incurred_on), "category": row.get_category_display(), "label": row.label, "amount": str(row.amount), "currency": row.currency.code} for row in expenses.exclude(category="SALARY")]
    salary_expenses = expenses.filter(category="SALARY").select_related("agent")
    if agent: salary_expenses = salary_expenses.filter(agent=agent)
    salary_rows = [{"agent": row.agent.get_full_name() or row.agent.email, "email": row.agent.email, "date": str(row.incurred_on), "amount": str(row.amount), "currency": row.currency.code} for row in salary_expenses if row.agent]
    periods = ProfitPeriod.objects.filter(start_date__lte=end, end_date__gte=start, status="FINALIZED")
    distributions = Distribution.objects.select_related("allocation__period__currency", "stakeholder").filter(allocation__period__in=periods)
    investments = Investment.objects.select_related("stakeholder", "currency").filter(invested_on__lte=end)
    partner_rows = PartnerOperation.objects.select_related("stakeholder", "operation__currency").filter(
        operation__status=OperationStatus.COMPLETED, operation__created_at__gte=start_at, operation__created_at__lt=end_at)
    stakeholders = Stakeholder.objects.select_related("owner").prefetch_related(
        Prefetch("contracts", queryset=Contract.objects.order_by("-starts_on", "-id"), to_attr="ordered_contracts")
    ).filter(is_active=True)
    # Apports (FundContribution) reçus dans la période
    contributions_qs = FundContribution.objects.select_related(
        "stakeholder", "currency", "receiving_account", "batch"
    ).filter(
        batch__status__in=["POSTED", "REVERSED"],
        batch__effective_at__gte=start_at,
        batch__effective_at__lt=end_at,
    )
    if stakeholder_type:
        stakeholders = stakeholders.filter(type=stakeholder_type)
        distributions = distributions.filter(stakeholder__type=stakeholder_type)
        investments = investments.filter(stakeholder__type=stakeholder_type)
        partner_rows = partner_rows.filter(stakeholder__type=stakeholder_type)
        contributions_qs = contributions_qs.filter(stakeholder__type=stakeholder_type)
    if stakeholder:
        stakeholders = stakeholders.filter(pk=stakeholder.pk)
        distributions = distributions.filter(stakeholder=stakeholder)
        investments = investments.filter(stakeholder=stakeholder)
        partner_rows = partner_rows.filter(stakeholder=stakeholder)
        contributions_qs = contributions_qs.filter(stakeholder=stakeholder)
    distribution_rows = [{"party": row.stakeholder.name, "type": row.stakeholder.get_type_display(), "amount": str(row.amount), "currency": row.allocation.period.currency.code, "status": row.get_status_display()} for row in distributions]
    frequency_divisor = {PaymentFrequency.MONTHLY: Decimal("12"), PaymentFrequency.QUARTERLY: Decimal("4"), PaymentFrequency.SEMIANNUAL: Decimal("2"), PaymentFrequency.ANNUAL: Decimal("1"), PaymentFrequency.AT_MATURITY: Decimal("1")}
    investment_rows = []
    for row in investments:
        annual_return = (row.amount * row.stakeholder.investor_return_percent / Decimal("100")).quantize(Decimal("0.01"))
        installment = (annual_return / frequency_divisor[row.stakeholder.payment_frequency]).quantize(Decimal("0.01"))
        investment_rows.append({"party": row.stakeholder.name, "identifier": str(row.stakeholder.public_id), "amount": str(row.amount), "currency": row.currency.code, "date": str(row.invested_on), "return_percent": str(row.stakeholder.investor_return_percent), "frequency": row.stakeholder.get_payment_frequency_display(), "annual_payment": str(annual_return), "payment_per_due_date": str(installment)})
    contribution_rows = [{
        "date": c.received_on.strftime("%d/%m/%Y"),
        "origin": c.get_origin_display(),
        "party": c.stakeholder.name,
        "currency": c.currency.code,
        "amount": str(c.amount),
        "account": c.receiving_account.code,
        "reference": c.external_reference or "",
    } for c in contributions_qs.order_by("received_on", "id")]
    partner_commissions = [{"partner": row.stakeholder.name, "operation": str(row.operation.reference), "share_percent": str(row.share_percent), "share": str((row.operation.fee * row.share_percent / Decimal("100")).quantize(Decimal("0.01"))), "currency": row.operation.currency.code} for row in partner_rows]
    totals = {}
    # A cancelled operation's amount/fee were fully reversed, so only
    # COMPLETED operations count towards the operations/commissions totals
    # (the "Operations" sheet itself still lists every status for the record).
    completed_summary = operations.filter(status=OperationStatus.COMPLETED).values("currency__code").annotate(volume=Sum("amount"), commissions=Sum("fee"))
    for row in completed_summary:
        item = totals.setdefault(row["currency__code"], {"operations": Decimal("0"), "commissions": Decimal("0"), "expenses": Decimal("0"), "distributions": Decimal("0"), "contributions": Decimal("0")})
        item["operations"] += row["volume"] or Decimal("0")
        item["commissions"] += row["commissions"] or Decimal("0")
    for row in expense_rows:
        totals.setdefault(row["currency"], {"operations": Decimal("0"), "commissions": Decimal("0"), "expenses": Decimal("0"), "distributions": Decimal("0"), "contributions": Decimal("0")})["expenses"] += Decimal(row["amount"])
    for row in salary_rows:
        totals.setdefault(row["currency"], {"operations": Decimal("0"), "commissions": Decimal("0"), "expenses": Decimal("0"), "distributions": Decimal("0"), "contributions": Decimal("0")})["expenses"] += Decimal(row["amount"])
    for row in distribution_rows:
        totals.setdefault(row["currency"], {"operations": Decimal("0"), "commissions": Decimal("0"), "expenses": Decimal("0"), "distributions": Decimal("0"), "contributions": Decimal("0")})["distributions"] += Decimal(row["amount"])
    for row in contribution_rows:
        totals.setdefault(row["currency"], {"operations": Decimal("0"), "commissions": Decimal("0"), "expenses": Decimal("0"), "distributions": Decimal("0"), "contributions": Decimal("0")})["contributions"] += Decimal(row["amount"])
    stakeholder_rows = []
    for party in stakeholders.order_by("type", "name"):
        contract = party.ordered_contracts[0] if party.ordered_contracts else None
        shareholder_value = (party.share_count * party.share_unit_value).quantize(Decimal("0.01"))
        estimated_dividend = (shareholder_value * party.dividend_percent / Decimal("100")).quantize(Decimal("0.01"))
        stakeholder_rows.append({
            "identifier": str(party.public_id), "type": party.get_type_display(), "name": party.name,
            "email": party.email, "phone": party.owner.phone if party.owner else "", "city": party.owner.city if party.owner else "",
            "contract": contract.title if contract else "", "starts_on": str(contract.starts_on) if contract else "",
            "ends_on": str(contract.ends_on) if contract and contract.ends_on else "Sans échéance",
            "clauses": contract.clauses if contract else "", "contract_status": contract.get_status_display() if contract else "",
            "return_percent": str(party.investor_return_percent) if party.type == "INVESTOR" else "-",
            "payment_frequency": party.get_payment_frequency_display() if party.type == "INVESTOR" else "-",
            "shares": str(party.share_count) if party.type == "SHAREHOLDER" else "-",
            "share_unit_value": str(party.share_unit_value) if party.type == "SHAREHOLDER" else "-",
            "dividend_percent": str(party.dividend_percent) if party.type == "SHAREHOLDER" else "-",
            "estimated_dividend": str(estimated_dividend) if party.type == "SHAREHOLDER" else "-",
            "partner_share_percent": str(party.partner_share_percent) if party.type == "PARTNER" else "-",
        })
    type_label = dict(Stakeholder._meta.get_field("type").choices).get(stakeholder_type, "Toutes")
    return {
        "period": {"start": str(start), "end": str(end)},
        "filters": {"agent": agent.email if agent else "Tous", "stakeholder": stakeholder.name if stakeholder else "Tous", "stakeholder_type": type_label},
        "operations": operation_rows, "expenses": expense_rows, "salaries": salary_rows,
        "investments": investment_rows, "contributions": contribution_rows,
        "stakeholders": stakeholder_rows, "partner_commissions": partner_commissions,
        "distributions": distribution_rows,
        "totals": {code: {key: str(value) for key, value in values.items()} for code, values in totals.items()},
    }


def _monthly_csv(snapshot):
    stream = StringIO(newline="")
    writer = csv.writer(stream)
    sections = [
        ("PARTIES_PRENANTES", snapshot["stakeholders"]),
        ("APPORTS_RECUS", snapshot["contributions"]),
        ("OPERATIONS", snapshot["operations"]),
        ("DEPENSES", snapshot["expenses"]),
        ("SALAIRES", snapshot["salaries"]),
        ("INVESTISSEMENTS", snapshot["investments"]),
        ("COMMISSIONS_PARTENAIRES", snapshot["partner_commissions"]),
        ("DIVIDENDES_ET_PAIES", snapshot["distributions"]),
    ]
    # Totals summary at the top
    writer.writerow(["SYNTHESE_PAR_DEVISE"])
    writer.writerow(["Devise", "Operations", "Commissions", "Depenses", "Distributions", "Apports"])
    for code, values in snapshot["totals"].items():
        writer.writerow(_safe_row([code, values["operations"], values["commissions"], values["expenses"], values["distributions"], values.get("contributions", "0.00")]))
    writer.writerow([])
    for title, rows in sections:
        writer.writerow([title])
        if rows:
            writer.writerow(rows[0].keys())
            for row in rows:
                writer.writerow(_safe_row(row.values()))
        writer.writerow([])
    return ("\ufeff" + stream.getvalue()).encode("utf-8")


def _monthly_xlsx(snapshot):
    workbook = Workbook()
    workbook.remove(workbook.active)
    sections = [
        ("Parties prenantes", snapshot["stakeholders"]),
        ("Apports reçus", snapshot["contributions"]),
        ("Opérations", snapshot["operations"]),
        ("Dépenses", snapshot["expenses"]),
        ("Salaires", snapshot["salaries"]),
        ("Investissements", snapshot["investments"]),
        ("Partenaires", snapshot["partner_commissions"]),
        ("Distributions", snapshot["distributions"]),
    ]
    for title, rows in sections:
        sheet = workbook.create_sheet(title)
        if rows:
            sheet.append(list(rows[0].keys()))
            for row in rows:
                sheet.append(_safe_row(row.values()))
            sheet.freeze_panes = "A2"
            sheet.auto_filter.ref = sheet.dimensions
    summary = workbook.create_sheet("Synthèse", 0)
    summary.append(["Période", snapshot["period"]["start"], snapshot["period"]["end"]])
    summary.append(["Agent", snapshot["filters"]["agent"]])
    summary.append(["Catégorie", snapshot["filters"]["stakeholder_type"]])
    summary.append(["Partie prenante", snapshot["filters"]["stakeholder"]])
    summary.append([])
    summary.append(["Devise", "Opérations", "Commissions", "Dépenses", "Distributions", "Apports reçus"])
    for code, values in snapshot["totals"].items():
        summary.append([code, values["operations"], values["commissions"], values["expenses"], values["distributions"], values.get("contributions", "0.00")])
    _style_workbook(workbook)
    stream = BytesIO()
    workbook.save(stream)
    return stream.getvalue()


def _fmt_money(val, currency=None):
    if val is None or str(val).strip() in ("", "-"):
        return "-"
    if str(val).strip().upper() == "N/A":
        return "N/A"
    try:
        d = Decimal(str(val))
        s = f"{d:,.2f}".replace(",", " ").replace(".", ",")
        return f"{s} {currency}" if currency else s
    except Exception:
        return str(val)


def _fmt_percent(val):
    if val is None or str(val).strip() in ("", "-"):
        return "-"
    if str(val).strip().upper() == "N/A":
        return "N/A"
    try:
        d = Decimal(str(val))
        s = f"{d:,.2f}".replace(",", " ").replace(".", ",")
        return f"{s} %"
    except Exception:
        return f"{val} %"


def _fmt_date(val):
    if not val or str(val).strip() in ("", "-"):
        return "-"
    s = str(val).strip()
    if s == "Sans échéance":
        return s
    if len(s) == 10 and s[4] == "-" and s[7] == "-":
        parts = s.split("-")
        return f"{parts[2]}/{parts[1]}/{parts[0]}"
    return s


def _escape_xml(text):
    if text is None:
        return ""
    return str(text).replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")


class _MonthlyNumberedCanvas(canvas.Canvas):
    """Canvas that computes total pages and draws headers & footers dynamically."""
    def __init__(self, *args, **kwargs):
        super().__init__(*args, **kwargs)
        self._saved_page_states = []
        self.doc_title = kwargs.get("doc_title", "RAPPORT FINANCIER MENSUEL")
        self.downloader = kwargs.get("downloader", "FINCORYA Group")
        self.logo_path = kwargs.get("logo_path", None)

    def showPage(self):
        self._saved_page_states.append(dict(self.__dict__))
        self._startPage()

    def save(self):
        num_pages = len(self._saved_page_states)
        for state in self._saved_page_states:
            self.__dict__.update(state)
            self.draw_page_elements(num_pages)
            super().showPage()
        super().save()

    def draw_page_elements(self, page_count):
        self.saveState()
        page_w, page_h = self._pagesize
        left_m = 12 * mm
        right_m = page_w - 12 * mm
        regular_font, bold_font = _report_fonts()

        # Header (Pages >= 2)
        if self._pageNumber > 1:
            self.setFont(bold_font, 8)
            self.setFillColor(colors.HexColor("#013B36"))
            if self.logo_path and Path(self.logo_path).exists():
                self.drawImage(
                    str(self.logo_path), left_m, page_h - 15 * mm,
                    width=26 * mm, height=9 * mm, preserveAspectRatio=True, mask="auto"
                )
            else:
                self.drawString(left_m, page_h - 12 * mm, "FINCORYA GROUP")

            self.drawRightString(right_m, page_h - 12 * mm, "RAPPORT FINANCIER MENSUEL")

            # Gold separator line
            self.setStrokeColor(colors.HexColor("#C9A227"))
            self.setLineWidth(0.75)
            self.line(left_m, page_h - 17 * mm, right_m, page_h - 17 * mm)

        # Footer (All pages)
        self.setStrokeColor(colors.HexColor("#DCE3DF"))
        self.setLineWidth(0.5)
        self.line(left_m, 11 * mm, right_m, 11 * mm)

        self.setFont(regular_font, 7.5)
        self.setFillColor(colors.HexColor("#596763"))
        self.drawString(left_m, 6.5 * mm, "FINCORYA Group  —  Document confidentiel")
        self.drawCentredString(page_w / 2, 6.5 * mm, f"Téléchargé par : {self.downloader}")
        self.drawRightString(right_m, 6.5 * mm, f"Page {self._pageNumber} sur {page_count}")

        self.restoreState()


def _monthly_pdf(snapshot, user):
    """
    Builds the executive-grade monthly financial PDF report according
    to the FINCORYA design specification and brand guidelines.
    """
    stream = BytesIO()
    regular_font, bold_font = _report_fonts()
    page_w, page_h = landscape(A4)
    left_m = 12 * mm
    right_m = 12 * mm
    top_m = 19 * mm
    bottom_m = 15 * mm
    usable_width = page_w - left_m - right_m

    doc = SimpleDocTemplate(
        stream,
        pagesize=landscape(A4),
        leftMargin=left_m,
        rightMargin=right_m,
        topMargin=top_m,
        bottomMargin=bottom_m,
        title="Rapport financier mensuel FINCORYA",
        author="FINCORYA Group",
    )

    downloader_name = user.get_full_name() or user.email if user else "Direction Générale"
    logo_path = _transparent_report_logo()

    # Brand Colors
    c_primary = colors.HexColor("#006B4F")
    c_dark = colors.HexColor("#013B36")
    c_gold = colors.HexColor("#C9A227")
    c_light = colors.HexColor("#F7F8F5")
    c_text = colors.HexColor("#18201F")
    c_muted = colors.HexColor("#596763")
    c_border = colors.HexColor("#DCE3DF")
    c_white = colors.HexColor("#FFFFFF")
    c_row_alt = colors.HexColor("#F9FAF8")

    # Typography Styles
    styles = {
        "Title": ParagraphStyle("DocTitle", fontName=bold_font, fontSize=18, leading=21, textColor=c_dark, alignment=TA_LEFT),
        "FilterText": ParagraphStyle("DocFilter", fontName=regular_font, fontSize=8.5, leading=11, textColor=c_text, alignment=TA_LEFT),
        "SectionHeading": ParagraphStyle("SectionH", fontName=bold_font, fontSize=11, leading=14, textColor=c_primary, spaceBefore=3.5 * mm, spaceAfter=1.8 * mm, keepWithNext=True),
        "SubSectionHeading": ParagraphStyle("SubSectionH", fontName=bold_font, fontSize=9, leading=11.5, textColor=c_dark, spaceBefore=2 * mm, spaceAfter=1.2 * mm, keepWithNext=True),
        "KpiLabel": ParagraphStyle("KpiL", fontName=bold_font, fontSize=7, leading=8.5, textColor=c_primary, alignment=TA_CENTER),
        "KpiValue": ParagraphStyle("KpiV", fontName=bold_font, fontSize=11.5, leading=13.5, textColor=c_dark, alignment=TA_CENTER),
        "KpiSub": ParagraphStyle("KpiS", fontName=regular_font, fontSize=6.5, leading=8, textColor=c_muted, alignment=TA_CENTER),
        "Th": ParagraphStyle("TableHead", fontName=bold_font, fontSize=7.5, leading=9, textColor=c_white, alignment=TA_LEFT),
        "ThRight": ParagraphStyle("TableHeadR", fontName=bold_font, fontSize=7.5, leading=9, textColor=c_white, alignment=TA_RIGHT),
        "ThCenter": ParagraphStyle("TableHeadC", fontName=bold_font, fontSize=7.5, leading=9, textColor=c_white, alignment=TA_CENTER),
        "Td": ParagraphStyle("TableData", fontName=regular_font, fontSize=7, leading=8.5, textColor=c_text, alignment=TA_LEFT),
        "TdBold": ParagraphStyle("TableDataB", fontName=bold_font, fontSize=7, leading=8.5, textColor=c_text, alignment=TA_LEFT),
        "TdRight": ParagraphStyle("TableDataR", fontName=regular_font, fontSize=7, leading=8.5, textColor=c_text, alignment=TA_RIGHT),
        "TdCenter": ParagraphStyle("TableDataC", fontName=regular_font, fontSize=7, leading=8.5, textColor=c_text, alignment=TA_CENTER),
        "TotalLabel": ParagraphStyle("TotLabel", fontName=bold_font, fontSize=7.5, leading=9, textColor=c_white, alignment=TA_LEFT),
        "TotalValue": ParagraphStyle("TotValue", fontName=bold_font, fontSize=7.5, leading=9, textColor=c_white, alignment=TA_RIGHT),
        "TotalCenter": ParagraphStyle("TotCenter", fontName=bold_font, fontSize=7.5, leading=9, textColor=c_white, alignment=TA_CENTER),
        "CardHead": ParagraphStyle("CardH", fontName=bold_font, fontSize=8, leading=10, textColor=c_white, alignment=TA_LEFT),
        "CardHeadRight": ParagraphStyle("CardHR", fontName=bold_font, fontSize=7.5, leading=9.5, textColor=c_gold, alignment=TA_RIGHT),
        "CardFieldValue": ParagraphStyle("CardFV", fontName=regular_font, fontSize=7, leading=8.5, textColor=c_text, alignment=TA_LEFT),
        "NoticeEmpty": ParagraphStyle("NoticeE", fontName=regular_font, fontSize=7.5, leading=9.5, textColor=c_muted, alignment=TA_LEFT),
    }

    story = []

    # =========================================================================
    # EN-TÊTE DE LA PREMIÈRE PAGE (Logo + Titre + Période + Filtres)
    # =========================================================================
    start_date_fr = _fmt_date(snapshot["period"]["start"])
    end_date_fr = _fmt_date(snapshot["period"]["end"])
    filter_agent = _escape_xml(snapshot["filters"].get("agent", "Tous"))
    filter_stakeholder = _escape_xml(snapshot["filters"].get("stakeholder", "Tous"))
    filter_category = _escape_xml(snapshot["filters"].get("stakeholder_type", "Toutes"))

    logo_element = ""
    if logo_path and Path(logo_path).exists():
        logo_element = Image(str(logo_path), width=35 * mm, height=14 * mm)
    else:
        logo_element = Paragraph("<b>FINCORYA GROUP</b>", styles["Title"])

    header_text_cells = [
        Paragraph("RAPPORT FINANCIER MENSUEL", styles["Title"]),
        Spacer(1, 0.8 * mm),
        Paragraph(
            f"<b>Période d'activité :</b> du {start_date_fr} au {end_date_fr}  &nbsp;|&nbsp;  "
            f"<b>Catégorie :</b> {filter_category}  &nbsp;|&nbsp;  "
            f"<b>Partie prenante :</b> {filter_stakeholder}  &nbsp;|&nbsp;  "
            f"<b>Agent :</b> {filter_agent}",
            styles["FilterText"]
        ),
    ]

    header_table = Table([[logo_element, header_text_cells]], colWidths=[42 * mm, usable_width - 42 * mm], hAlign="LEFT")
    header_table.setStyle(TableStyle([
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("LEFTPADDING", (0, 0), (-1, -1), 0),
        ("RIGHTPADDING", (0, 0), (-1, -1), 0),
        ("TOPPADDING", (0, 0), (-1, -1), 0),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 2),
    ]))
    story.append(header_table)
    story.append(Spacer(1, 1 * mm))
    story.append(HRFlowable(width="100%", thickness=1, color=c_gold, spaceBefore=0.5, spaceAfter=2.5 * mm))

    # =========================================================================
    # BLOCS CHIFFRÉS SOBRES (KPIs)
    # =========================================================================
    tot_ops = sum(Decimal(v["operations"]) for v in snapshot["totals"].values())
    tot_comm = sum(Decimal(v["commissions"]) for v in snapshot["totals"].values())
    tot_exp = sum(Decimal(v["expenses"]) for v in snapshot["totals"].values())
    tot_dist = sum(Decimal(v["distributions"]) for v in snapshot["totals"].values())
    tot_contrib = sum(Decimal(v.get("contributions", "0")) for v in snapshot["totals"].values())
    curr_label = list(snapshot["totals"].keys())[0] if len(snapshot["totals"]) == 1 else "USD"
    nb_ops = len(snapshot.get("operations", []))

    kpi_cards = [
        [
            Paragraph("VOLUME DES OPÉRATIONS", styles["KpiLabel"]),
            Paragraph(f"{_fmt_money(tot_ops)} <font size=7.5>{curr_label}</font>", styles["KpiValue"]),
            Paragraph(f"{nb_ops} opérations traitées", styles["KpiSub"]),
        ],
        [
            Paragraph("COMMISSIONS NETTES", styles["KpiLabel"]),
            Paragraph(f"{_fmt_money(tot_comm)} <font size=7.5>{curr_label}</font>", styles["KpiValue"]),
            Paragraph("Marge brute de services", styles["KpiSub"]),
        ],
        [
            Paragraph("DÉPENSES & CHARGES", styles["KpiLabel"]),
            Paragraph(f"{_fmt_money(tot_exp)} <font size=7.5>{curr_label}</font>", styles["KpiValue"]),
            Paragraph(f"{len(snapshot.get('expenses', []))} charges approuvées", styles["KpiSub"]),
        ],
        [
            Paragraph("DISTRIBUTIONS EFFECTUÉES", styles["KpiLabel"]),
            Paragraph(f"{_fmt_money(tot_dist)} <font size=7.5>{curr_label}</font>", styles["KpiValue"]),
            Paragraph(f"{len(snapshot.get('distributions', []))} allocations", styles["KpiSub"]),
        ],
        [
            Paragraph("APPORTS EN CAPITAL", styles["KpiLabel"]),
            Paragraph(f"{_fmt_money(tot_contrib)} <font size=7.5>{curr_label}</font>", styles["KpiValue"]),
            Paragraph(f"{len(snapshot.get('contributions', []))} versements reçus", styles["KpiSub"]),
        ],
    ]

    card_w = usable_width / 5.0
    kpi_table = Table([kpi_cards], colWidths=[card_w] * 5, hAlign="LEFT")
    kpi_table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, -1), c_light),
        ("GRID", (0, 0), (-1, -1), 0.5, c_border),
        ("LINEABOVE", (0, 0), (-1, 0), 2, c_primary),
        ("TOPPADDING", (0, 0), (-1, -1), 3),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
        ("LEFTPADDING", (0, 0), (-1, -1), 4),
        ("RIGHTPADDING", (0, 0), (-1, -1), 4),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
    ]))
    story.append(kpi_table)
    story.append(Spacer(1, 2.5 * mm))

    # =========================================================================
    # 1. SYNTHÈSE FINANCIÈRE PAR DEVISE
    # =========================================================================
    story.append(Paragraph("1. Synthèse financière par devise", styles["SectionHeading"]))

    synth_headers = [
        Paragraph("Devise", styles["ThCenter"]),
        Paragraph("Volume des opérations", styles["ThRight"]),
        Paragraph("Commissions nettes", styles["ThRight"]),
        Paragraph("Dépenses & charges", styles["ThRight"]),
        Paragraph("Distributions", styles["ThRight"]),
        Paragraph("Apports reçus", styles["ThRight"]),
    ]
    synth_data = [synth_headers]

    for curr_code, vals in sorted(snapshot["totals"].items()):
        synth_data.append([
            Paragraph(curr_code, styles["TdCenter"]),
            Paragraph(_fmt_money(vals["operations"]), styles["TdRight"]),
            Paragraph(_fmt_money(vals["commissions"]), styles["TdRight"]),
            Paragraph(_fmt_money(vals["expenses"]), styles["TdRight"]),
            Paragraph(_fmt_money(vals["distributions"]), styles["TdRight"]),
            Paragraph(_fmt_money(vals.get("contributions", "0")), styles["TdRight"]),
        ])

    synth_data.append([
        Paragraph("TOTAL CONSOLIDÉ", styles["TotalLabel"]),
        Paragraph(_fmt_money(tot_ops), styles["TotalValue"]),
        Paragraph(_fmt_money(tot_comm), styles["TotalValue"]),
        Paragraph(_fmt_money(tot_exp), styles["TotalValue"]),
        Paragraph(_fmt_money(tot_dist), styles["TotalValue"]),
        Paragraph(_fmt_money(tot_contrib), styles["TotalValue"]),
    ])

    col_dev = 54
    col_other = (usable_width - col_dev) / 5.0
    synth_widths = [col_dev, col_other, col_other, col_other, col_other, col_other]
    synth_table = Table(synth_data, colWidths=synth_widths, repeatRows=1, hAlign="LEFT")
    synth_table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), c_primary),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("GRID", (0, 0), (-1, -2), 0.5, c_border),
        ("ROWBACKGROUNDS", (0, 1), (-1, -2), [c_white, c_row_alt]),
        ("TOPPADDING", (0, 0), (-1, -1), 3),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
        ("LEFTPADDING", (0, 0), (-1, -1), 5),
        ("RIGHTPADDING", (0, 0), (-1, -1), 5),
        ("BACKGROUND", (0, -1), (-1, -1), c_dark),
        ("LINEABOVE", (0, -1), (-1, -1), 1, c_gold),
    ]))
    story.append(synth_table)
    story.append(Spacer(1, 3 * mm))

    # =========================================================================
    # 2. DÉPENSES ET CHARGES
    # =========================================================================
    story.append(Paragraph("2. Dépenses et charges", styles["SectionHeading"]))
    expenses = snapshot.get("expenses", [])

    if expenses:
        exp_headers = [
            Paragraph("Date", styles["ThCenter"]),
            Paragraph("Catégorie", styles["Th"]),
            Paragraph("Libellé / Justificatif", styles["Th"]),
            Paragraph("Montant", styles["ThRight"]),
            Paragraph("Devise", styles["ThCenter"]),
        ]
        exp_data = [exp_headers]
        tot_exp_cat = Decimal("0")
        exp_curr = "USD"
        for r in expenses:
            amt = Decimal(r["amount"])
            tot_exp_cat += amt
            exp_curr = r["currency"]
            exp_data.append([
                Paragraph(_fmt_date(r["date"]), styles["TdCenter"]),
                Paragraph(_escape_xml(r["category"]), styles["TdBold"]),
                Paragraph(_escape_xml(r["label"]), styles["Td"]),
                Paragraph(_fmt_money(amt), styles["TdRight"]),
                Paragraph(r["currency"], styles["TdCenter"]),
            ])
        exp_data.append([
            Paragraph("TOTAL DÉPENSES & CHARGES", styles["TotalLabel"]),
            Paragraph("", styles["TotalLabel"]),
            Paragraph(f"{len(expenses)} enregistrements validés", styles["TotalLabel"]),
            Paragraph(_fmt_money(tot_exp_cat), styles["TotalValue"]),
            Paragraph(exp_curr, styles["TotalCenter"]),
        ])
        exp_widths = [65, 125, usable_width - 65 - 125 - 85 - 45, 85, 45]
        exp_table = Table(exp_data, colWidths=exp_widths, repeatRows=1, hAlign="LEFT")
        exp_table.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), c_primary),
            ("GRID", (0, 0), (-1, -2), 0.5, c_border),
            ("ROWBACKGROUNDS", (0, 1), (-1, -2), [c_white, c_row_alt]),
            ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
            ("TOPPADDING", (0, 0), (-1, -1), 2.5),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 2.5),
            ("LEFTPADDING", (0, 0), (-1, -1), 4),
            ("RIGHTPADDING", (0, 0), (-1, -1), 4),
            ("BACKGROUND", (0, -1), (-1, -1), c_dark),
            ("LINEABOVE", (0, -1), (-1, -1), 1, c_gold),
            ("SPAN", (0, -1), (1, -1)),
        ]))
        story.append(exp_table)
    else:
        empty_box = Table(
            [[Paragraph("<i>Aucune dépense ni charge enregistrée pour cette période.</i>", styles["NoticeEmpty"])]],
            colWidths=[usable_width], hAlign="LEFT"
        )
        empty_box.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, -1), c_light),
            ("BOX", (0, 0), (-1, -1), 0.5, c_border),
            ("PADDING", (0, 0), (-1, -1), 4),
        ]))
        story.append(empty_box)

    story.append(Spacer(1, 3 * mm))

    # =========================================================================
    # 3. SALAIRES DES AGENTS
    # =========================================================================
    story.append(Paragraph("3. Salaires des agents", styles["SectionHeading"]))
    salaries = snapshot.get("salaries", [])

    if salaries:
        sal_headers = [
            Paragraph("Agent", styles["Th"]),
            Paragraph("E-mail", styles["Th"]),
            Paragraph("Date", styles["ThCenter"]),
            Paragraph("Montant", styles["ThRight"]),
            Paragraph("Devise", styles["ThCenter"]),
        ]
        sal_data = [sal_headers]
        tot_sal = Decimal("0")
        sal_curr = "USD"
        for r in salaries:
            amt = Decimal(r["amount"])
            tot_sal += amt
            sal_curr = r["currency"]
            sal_data.append([
                Paragraph(_escape_xml(r["agent"]), styles["TdBold"]),
                Paragraph(_escape_xml(r["email"]), styles["Td"]),
                Paragraph(_fmt_date(r["date"]), styles["TdCenter"]),
                Paragraph(_fmt_money(amt), styles["TdRight"]),
                Paragraph(r["currency"], styles["TdCenter"]),
            ])
        sal_data.append([
            Paragraph("TOTAL SALAIRES", styles["TotalLabel"]),
            Paragraph(f"{len(salaries)} versements", styles["TotalLabel"]),
            Paragraph("", styles["TotalLabel"]),
            Paragraph(_fmt_money(tot_sal), styles["TotalValue"]),
            Paragraph(sal_curr, styles["TotalCenter"]),
        ])
        sal_widths = [140, 180, 80, 85, usable_width - 140 - 180 - 80 - 85]
        sal_table = Table(sal_data, colWidths=sal_widths, repeatRows=1, hAlign="LEFT")
        sal_table.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), c_primary),
            ("GRID", (0, 0), (-1, -2), 0.5, c_border),
            ("ROWBACKGROUNDS", (0, 1), (-1, -2), [c_white, c_row_alt]),
            ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
            ("TOPPADDING", (0, 0), (-1, -1), 2.5),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 2.5),
            ("LEFTPADDING", (0, 0), (-1, -1), 4),
            ("RIGHTPADDING", (0, 0), (-1, -1), 4),
            ("BACKGROUND", (0, -1), (-1, -1), c_dark),
            ("LINEABOVE", (0, -1), (-1, -1), 1, c_gold),
        ]))
        story.append(sal_table)
    else:
        empty_box = Table(
            [[Paragraph("<i>Aucun salaire d'agent enregistré sur cette période.</i>", styles["NoticeEmpty"])]],
            colWidths=[usable_width], hAlign="LEFT"
        )
        empty_box.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, -1), c_light),
            ("BOX", (0, 0), (-1, -1), 0.5, c_border),
            ("PADDING", (0, 0), (-1, -1), 4),
        ]))
        story.append(empty_box)

    story.append(Spacer(1, 3 * mm))

    # =========================================================================
    # 4. INVESTISSEMENTS ET PAIEMENTS CALCULÉS (& APPORTS REÇUS)
    # =========================================================================
    story.append(PageBreak())
    story.append(Paragraph("4. Investissements et paiements calculés", styles["SectionHeading"]))
    investments = snapshot.get("investments", [])
    contributions = snapshot.get("contributions", [])

    uuid_map = {}
    for party in snapshot.get("stakeholders", []):
        uid = party["identifier"]
        uuid_map[uid] = f"[{uid[:8]}]"

    if investments:
        story.append(Paragraph("4.1. Investissements contractuels", styles["SubSectionHeading"]))
        inv_headers = [
            Paragraph("Partie prenante", styles["Th"]),
            Paragraph("Réf.", styles["ThCenter"]),
            Paragraph("Date", styles["ThCenter"]),
            Paragraph("Montant investi", styles["ThRight"]),
            Paragraph("Devise", styles["ThCenter"]),
            Paragraph("Rendement", styles["ThRight"]),
            Paragraph("Fréquence", styles["ThCenter"]),
            Paragraph("Paiement annuel", styles["ThRight"]),
            Paragraph("Paiement / échéance", styles["ThRight"]),
        ]
        inv_data = [inv_headers]
        tot_inv = Decimal("0")
        tot_ann = Decimal("0")
        tot_inst = Decimal("0")
        inv_curr = "USD"
        for r in investments:
            tot_inv += Decimal(r["amount"])
            tot_ann += Decimal(r["annual_payment"])
            tot_inst += Decimal(r["payment_per_due_date"])
            inv_curr = r["currency"]
            uid_short = uuid_map.get(r["identifier"], f"[{r['identifier'][:8]}]")
            inv_data.append([
                Paragraph(_escape_xml(r["party"]), styles["TdBold"]),
                Paragraph(uid_short, styles["TdCenter"]),
                Paragraph(_fmt_date(r["date"]), styles["TdCenter"]),
                Paragraph(_fmt_money(r["amount"]), styles["TdRight"]),
                Paragraph(r["currency"], styles["TdCenter"]),
                Paragraph(_fmt_percent(r["return_percent"]), styles["TdRight"]),
                Paragraph(_escape_xml(r["frequency"]), styles["TdCenter"]),
                Paragraph(_fmt_money(r["annual_payment"]), styles["TdRight"]),
                Paragraph(_fmt_money(r["payment_per_due_date"]), styles["TdRight"]),
            ])
        inv_data.append([
            Paragraph("TOTAL INVESTISSEMENTS", styles["TotalLabel"]),
            Paragraph("", styles["TotalLabel"]),
            Paragraph("", styles["TotalLabel"]),
            Paragraph(_fmt_money(tot_inv), styles["TotalValue"]),
            Paragraph(inv_curr, styles["TotalCenter"]),
            Paragraph("-", styles["TotalCenter"]),
            Paragraph("-", styles["TotalCenter"]),
            Paragraph(_fmt_money(tot_ann), styles["TotalValue"]),
            Paragraph(_fmt_money(tot_inst), styles["TotalValue"]),
        ])
        inv_widths = [135, 60, 60, 75, 40, 65, 80, 85, 85]
        scale = usable_width / sum(inv_widths)
        inv_widths = [w * scale for w in inv_widths]
        inv_table = Table(inv_data, colWidths=inv_widths, repeatRows=1, hAlign="LEFT")
        inv_table.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), c_primary),
            ("GRID", (0, 0), (-1, -2), 0.5, c_border),
            ("ROWBACKGROUNDS", (0, 1), (-1, -2), [c_white, c_row_alt]),
            ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
            ("TOPPADDING", (0, 0), (-1, -1), 2.5),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 2.5),
            ("LEFTPADDING", (0, 0), (-1, -1), 4),
            ("RIGHTPADDING", (0, 0), (-1, -1), 4),
            ("BACKGROUND", (0, -1), (-1, -1), c_dark),
            ("LINEABOVE", (0, -1), (-1, -1), 1, c_gold),
            ("SPAN", (0, -1), (2, -1)),
        ]))
        story.append(inv_table)
        story.append(Spacer(1, 2 * mm))

    # Sub-table 4.2 Apports reçus
    if contributions:
        story.append(Paragraph("4.2. Apports reçus en capital et garanties", styles["SubSectionHeading"]))
        contrib_headers = [
            Paragraph("Date", styles["ThCenter"]),
            Paragraph("Origine de l'apport", styles["Th"]),
            Paragraph("Partie prenante", styles["Th"]),
            Paragraph("Compte de réception", styles["Th"]),
            Paragraph("Référence externe", styles["Th"]),
            Paragraph("Montant", styles["ThRight"]),
            Paragraph("Devise", styles["ThCenter"]),
        ]
        contrib_data = [contrib_headers]
        tot_contrib_sec = Decimal("0")
        contrib_curr = "USD"
        for r in contributions:
            amt = Decimal(r["amount"])
            tot_contrib_sec += amt
            contrib_curr = r["currency"]
            contrib_data.append([
                Paragraph(_fmt_date(r["date"]), styles["TdCenter"]),
                Paragraph(_escape_xml(r["origin"]), styles["TdBold"]),
                Paragraph(_escape_xml(r["party"]), styles["Td"]),
                Paragraph(_escape_xml(r["account"]), styles["Td"]),
                Paragraph(_escape_xml(r.get("reference", "-")), styles["Td"]),
                Paragraph(_fmt_money(amt), styles["TdRight"]),
                Paragraph(r["currency"], styles["TdCenter"]),
            ])
        contrib_data.append([
            Paragraph("TOTAL DES APPORTS REÇUS", styles["TotalLabel"]),
            Paragraph("", styles["TotalLabel"]),
            Paragraph(f"{len(contributions)} apports enregistrés", styles["TotalLabel"]),
            Paragraph("", styles["TotalLabel"]),
            Paragraph("", styles["TotalLabel"]),
            Paragraph(_fmt_money(tot_contrib_sec), styles["TotalValue"]),
            Paragraph(contrib_curr, styles["TotalCenter"]),
        ])
        c_widths = [65, 120, 130, 120, 110, 85, 45]
        scale = usable_width / sum(c_widths)
        c_widths = [w * scale for w in c_widths]
        contrib_table = Table(contrib_data, colWidths=c_widths, repeatRows=1, hAlign="LEFT")
        contrib_table.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), c_primary),
            ("GRID", (0, 0), (-1, -2), 0.5, c_border),
            ("ROWBACKGROUNDS", (0, 1), (-1, -2), [c_white, c_row_alt]),
            ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
            ("TOPPADDING", (0, 0), (-1, -1), 2.5),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 2.5),
            ("LEFTPADDING", (0, 0), (-1, -1), 4),
            ("RIGHTPADDING", (0, 0), (-1, -1), 4),
            ("BACKGROUND", (0, -1), (-1, -1), c_dark),
            ("LINEABOVE", (0, -1), (-1, -1), 1, c_gold),
            ("SPAN", (0, -1), (1, -1)),
        ]))
        story.append(contrib_table)

    story.append(Spacer(1, 3 * mm))

    # =========================================================================
    # 5. COMMISSIONS PARTENAIRES & 6. DIVIDENDES ET PAIEMENTS
    # =========================================================================
    story.append(Paragraph("5. Commissions partenaires", styles["SectionHeading"]))
    partner_comms = snapshot.get("partner_commissions", [])

    if partner_comms:
        pc_headers = [
            Paragraph("Partenaire", styles["Th"]),
            Paragraph("Opération réf.", styles["ThCenter"]),
            Paragraph("Part contractuelle", styles["ThRight"]),
            Paragraph("Commission partenaire", styles["ThRight"]),
            Paragraph("Devise", styles["ThCenter"]),
        ]
        pc_data = [pc_headers]
        tot_pc = Decimal("0")
        pc_curr = "USD"
        for r in partner_comms:
            amt = Decimal(r["share"])
            tot_pc += amt
            pc_curr = r["currency"]
            op_ref_raw = r["operation"]
            op_short = f"[{op_ref_raw[:8]}]" if len(op_ref_raw) > 12 else op_ref_raw
            pc_data.append([
                Paragraph(_escape_xml(r["partner"]), styles["TdBold"]),
                Paragraph(op_short, styles["TdCenter"]),
                Paragraph(_fmt_percent(r["share_percent"]), styles["TdRight"]),
                Paragraph(_fmt_money(amt), styles["TdRight"]),
                Paragraph(r["currency"], styles["TdCenter"]),
            ])
        pc_data.append([
            Paragraph("TOTAL COMMISSIONS PARTENAIRES", styles["TotalLabel"]),
            Paragraph(f"{len(partner_comms)} opérations", styles["TotalCenter"]),
            Paragraph("", styles["TotalLabel"]),
            Paragraph(_fmt_money(tot_pc), styles["TotalValue"]),
            Paragraph(pc_curr, styles["TotalCenter"]),
        ])
        pc_widths = [200, 140, 100, 100, usable_width - 200 - 140 - 100 - 100]
        pc_table = Table(pc_data, colWidths=pc_widths, repeatRows=1, hAlign="LEFT")
        pc_table.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), c_primary),
            ("GRID", (0, 0), (-1, -2), 0.5, c_border),
            ("ROWBACKGROUNDS", (0, 1), (-1, -2), [c_white, c_row_alt]),
            ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
            ("TOPPADDING", (0, 0), (-1, -1), 2.5),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 2.5),
            ("LEFTPADDING", (0, 0), (-1, -1), 4),
            ("RIGHTPADDING", (0, 0), (-1, -1), 4),
            ("BACKGROUND", (0, -1), (-1, -1), c_dark),
            ("LINEABOVE", (0, -1), (-1, -1), 1, c_gold),
        ]))
        story.append(pc_table)
    else:
        empty_box = Table(
            [[Paragraph("<i>Aucune commission partenaire enregistrée sur cette période.</i>", styles["NoticeEmpty"])]],
            colWidths=[usable_width], hAlign="LEFT"
        )
        empty_box.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, -1), c_light),
            ("BOX", (0, 0), (-1, -1), 0.5, c_border),
            ("PADDING", (0, 0), (-1, -1), 4),
        ]))
        story.append(empty_box)

    story.append(Spacer(1, 3 * mm))

    # =========================================================================
    # 6. DIVIDENDES ET PAIEMENTS
    # =========================================================================
    story.append(Paragraph("6. Dividendes et paiements", styles["SectionHeading"]))
    distributions = snapshot.get("distributions", [])

    if distributions:
        dist_headers = [
            Paragraph("Bénéficiaire", styles["Th"]),
            Paragraph("Type", styles["ThCenter"]),
            Paragraph("Montant alloué", styles["ThRight"]),
            Paragraph("Devise", styles["ThCenter"]),
            Paragraph("Statut", styles["ThCenter"]),
        ]
        dist_data = [dist_headers]
        tot_dist_sec = Decimal("0")
        dist_curr = "USD"
        for r in distributions:
            amt = Decimal(r["amount"])
            tot_dist_sec += amt
            dist_curr = r["currency"]
            dist_data.append([
                Paragraph(_escape_xml(r["party"]), styles["TdBold"]),
                Paragraph(_escape_xml(r["type"]), styles["TdCenter"]),
                Paragraph(_fmt_money(amt), styles["TdRight"]),
                Paragraph(r["currency"], styles["TdCenter"]),
                Paragraph(_escape_xml(r["status"]), styles["TdCenter"]),
            ])
        dist_data.append([
            Paragraph("TOTAL DISTRIBUTIONS EFFECTUÉES", styles["TotalLabel"]),
            Paragraph(f"{len(distributions)} distributions", styles["TotalCenter"]),
            Paragraph(_fmt_money(tot_dist_sec), styles["TotalValue"]),
            Paragraph(dist_curr, styles["TotalCenter"]),
            Paragraph("", styles["TotalCenter"]),
        ])
        dist_widths = [240, 140, 120, 50, usable_width - 240 - 140 - 120 - 50]
        dist_table = Table(dist_data, colWidths=dist_widths, repeatRows=1, hAlign="LEFT")
        dist_table.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), c_primary),
            ("GRID", (0, 0), (-1, -2), 0.5, c_border),
            ("ROWBACKGROUNDS", (0, 1), (-1, -2), [c_white, c_row_alt]),
            ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
            ("TOPPADDING", (0, 0), (-1, -1), 2.5),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 2.5),
            ("LEFTPADDING", (0, 0), (-1, -1), 4),
            ("RIGHTPADDING", (0, 0), (-1, -1), 4),
            ("BACKGROUND", (0, -1), (-1, -1), c_dark),
            ("LINEABOVE", (0, -1), (-1, -1), 1, c_gold),
        ]))
        story.append(dist_table)
    else:
        empty_box = Table(
            [[Paragraph("<i>Aucune distribution de dividende enregistrée sur cette période.</i>", styles["NoticeEmpty"])]],
            colWidths=[usable_width], hAlign="LEFT"
        )
        empty_box.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, -1), c_light),
            ("BOX", (0, 0), (-1, -1), 0.5, c_border),
            ("PADDING", (0, 0), (-1, -1), 4),
        ]))
        story.append(empty_box)

    story.append(Spacer(1, 3 * mm))

    # =========================================================================
    # 7. FICHES DES PARTIES PRENANTES
    # =========================================================================
    story.append(PageBreak())
    story.append(Paragraph("7. Fiches des parties prenantes", styles["SectionHeading"]))
    stakeholders = snapshot.get("stakeholders", [])

    half_card_w = (usable_width - 4 * mm) / 2.0

    cards_flowables = []
    for party in stakeholders:
        uid_short = f"[{party['identifier'][:8]}]"
        ptype = party.get("type", "Partie prenante")
        pname = party.get("name", "")
        pemail = party.get("email") or "Non renseigné"
        pphone = party.get("phone") or "Non renseigné"
        pcity = party.get("city") or "Non renseignée"
        pcontract = party.get("contract") or "Sans contrat"
        pstarts = _fmt_date(party.get("starts_on"))
        pends = _fmt_date(party.get("ends_on"))
        pclauses = party.get("clauses") or "Aucune clause enregistrée"
        pstatus = party.get("contract_status") or "Actif"

        if ptype == "Investisseur":
            at_maturity_label = "À l'échéance"
            fin_info = (
                f"<b>Rendement contractuel :</b> {_fmt_percent(party.get('return_percent'))}  &nbsp;|&nbsp;  "
                f"<b>Fréquence de paiement :</b> {party.get('payment_frequency') or at_maturity_label}<br/>"
                f"<b>Parts sociales :</b> <font color='#596763'>N/A</font>  &nbsp;|&nbsp;  "
                f"<b>Dividende :</b> <font color='#596763'>N/A</font>  &nbsp;|&nbsp;  "
                f"<b>Part partenaire :</b> <font color='#596763'>N/A</font>"
            )
        elif ptype == "Actionnaire":
            shares_str = party.get("shares")
            unit_val_str = party.get("share_unit_value")
            tot_shares_val = "-"
            try:
                tot_shares_val = _fmt_money(Decimal(shares_str) * Decimal(unit_val_str), "USD")
            except Exception:
                tot_shares_val = "-"
            fin_info = (
                f"<b>Parts détenues :</b> {shares_str}  &nbsp;|&nbsp;  "
                f"<b>Valeur unitaire :</b> {_fmt_money(unit_val_str, 'USD')}  &nbsp;|&nbsp;  "
                f"<b>Capital :</b> {tot_shares_val}<br/>"
                f"<b>Dividende contractuel :</b> {_fmt_percent(party.get('dividend_percent'))}  &nbsp;|&nbsp;  "
                f"<b>Dividende estimé :</b> {_fmt_money(party.get('estimated_dividend'), 'USD')}<br/>"
                f"<b>Rendement investisseur :</b> <font color='#596763'>N/A</font>  &nbsp;|&nbsp;  "
                f"<b>Part partenaire :</b> <font color='#596763'>N/A</font>"
            )
        elif ptype == "Partenaire":
            fin_info = (
                f"<b>Part commissions :</b> {_fmt_percent(party.get('partner_share_percent'))}<br/>"
                f"<b>Rendement investisseur :</b> <font color='#596763'>N/A</font>  &nbsp;|&nbsp;  "
                f"<b>Parts sociales :</b> <font color='#596763'>N/A</font>  &nbsp;|&nbsp;  "
                f"<b>Dividende :</b> <font color='#596763'>N/A</font>"
            )
        else:
            fin_info = "<font color='#596763'>Aucune condition financière enregistrée</font>"

        card_rows = [
            [
                Paragraph(f"<b>{_escape_xml(pname)}</b> &nbsp;·&nbsp; <font size=7 color='#DCE3DF'>{_escape_xml(ptype)}</font>", styles["CardHead"]),
                Paragraph(f"Réf. {uid_short}", styles["CardHeadRight"]),
            ],
            [
                Paragraph(
                    f"<b>Identité & Coordonnées :</b> {_escape_xml(pemail)} &nbsp;·&nbsp; Tél : {_escape_xml(pphone)} &nbsp;·&nbsp; Ville : {_escape_xml(pcity)}",
                    styles["CardFieldValue"]
                ),
                "",
            ],
            [
                Paragraph(
                    f"<b>Contrat :</b> {_escape_xml(pcontract)} ({_escape_xml(pstatus)}) &nbsp;·&nbsp; "
                    f"Période : {pstarts} au {pends}<br/>"
                    f"<b>Clauses contractuelles :</b> {_escape_xml(pclauses)}",
                    styles["CardFieldValue"]
                ),
                "",
            ],
            [
                Paragraph(f"<b>Conditions financières :</b><br/>{fin_info}", styles["CardFieldValue"]),
                "",
            ],
        ]

        card_table = Table(card_rows, colWidths=[half_card_w - 60, 60], hAlign="LEFT")
        card_table.setStyle(TableStyle([
            ("SPAN", (0, 1), (1, 1)),
            ("SPAN", (0, 2), (1, 2)),
            ("SPAN", (0, 3), (1, 3)),
            ("BACKGROUND", (0, 0), (-1, 0), c_dark),
            ("BACKGROUND", (0, 1), (-1, -1), c_light),
            ("BOX", (0, 0), (-1, -1), 0.5, c_border),
            ("LINEBELOW", (0, 0), (-1, 0), 1, c_gold),
            ("LINEBELOW", (0, 1), (-1, -2), 0.35, c_border),
            ("TOPPADDING", (0, 0), (-1, -1), 2),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 2),
            ("LEFTPADDING", (0, 0), (-1, -1), 4),
            ("RIGHTPADDING", (0, 0), (-1, -1), 4),
            ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ]))
        cards_flowables.append(card_table)

    paired_cards = []
    for i in range(0, len(cards_flowables), 2):
        left_c = cards_flowables[i]
        right_c = cards_flowables[i + 1] if i + 1 < len(cards_flowables) else ""
        pair_row = Table([[left_c, right_c]], colWidths=[half_card_w, half_card_w], hAlign="LEFT")
        pair_row.setStyle(TableStyle([
            ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ("LEFTPADDING", (0, 0), (-1, -1), 0),
            ("RIGHTPADDING", (0, 0), (-1, -1), 2 * mm),
            ("TOPPADDING", (0, 0), (-1, -1), 0),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 2 * mm),
        ]))
        paired_cards.append(pair_row)

    for p in paired_cards:
        story.append(p)

    story.append(Spacer(1, 2 * mm))

    # Annexe de correspondance des identifiants (UUID)
    story.append(Paragraph("Annexe A : Table de correspondance des identifiants uniques des parties prenantes", styles["SubSectionHeading"]))
    annex_headers = [
        Paragraph("Repère", styles["ThCenter"]),
        Paragraph("Nom de la partie prenante", styles["Th"]),
        Paragraph("Rôle / Type", styles["Th"]),
        Paragraph("Identifiant unique complet (UUID)", styles["Th"]),
    ]
    annex_data = [annex_headers]
    for party in stakeholders:
        annex_data.append([
            Paragraph(f"[{party['identifier'][:8]}]", styles["TdCenter"]),
            Paragraph(_escape_xml(party.get("name")), styles["TdBold"]),
            Paragraph(_escape_xml(party.get("type")), styles["Td"]),
            Paragraph(f"<font name='Courier' size=6.5>{party['identifier']}</font>", styles["Td"]),
        ])
    annex_widths = [55, 160, 110, usable_width - 55 - 160 - 110]
    annex_table = Table(annex_data, colWidths=annex_widths, repeatRows=1, hAlign="LEFT")
    annex_table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), c_primary),
        ("GRID", (0, 0), (-1, -1), 0.5, c_border),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [c_white, c_row_alt]),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("TOPPADDING", (0, 0), (-1, -1), 2),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 2),
        ("LEFTPADDING", (0, 0), (-1, -1), 4),
        ("RIGHTPADDING", (0, 0), (-1, -1), 4),
    ]))
    story.append(annex_table)
    story.append(Spacer(1, 4 * mm))

    # =========================================================================
    # 8. JOURNAL DÉTAILLÉ DES OPÉRATIONS
    # =========================================================================
    story.append(PageBreak())
    story.append(Paragraph("8. Journal détaillé des opérations", styles["SectionHeading"]))
    operations = snapshot.get("operations", [])

    if operations:
        op_headers = [
            Paragraph("Date / Heure", styles["ThCenter"]),
            Paragraph("Réf.", styles["ThCenter"]),
            Paragraph("Type", styles["ThCenter"]),
            Paragraph("Service", styles["Th"]),
            Paragraph("Client", styles["Th"]),
            Paragraph("Identifiant client", styles["Th"]),
            Paragraph("Montant", styles["ThRight"]),
            Paragraph("Dev.", styles["ThCenter"]),
            Paragraph("Commission", styles["ThRight"]),
            Paragraph("Agent traitant", styles["Th"]),
        ]
        op_data = [op_headers]
        tot_op_amt = Decimal("0")
        tot_op_fees = Decimal("0")
        op_curr = "USD"

        for op in operations:
            amt = Decimal(op["amount"])
            fee = Decimal(op["commission"])
            op_curr = op["currency"]
            if op.get("status") in {"Terminée", "Complétée", "COMPLETED"}:
                tot_op_amt += amt
                tot_op_fees += fee

            ref_raw = op["reference"]
            ref_display = f"[{ref_raw[:8]}]" if len(ref_raw) > 12 else ref_raw

            op_data.append([
                Paragraph(op["date"], styles["TdCenter"]),
                Paragraph(ref_display, styles["TdCenter"]),
                Paragraph(_escape_xml(op["type"]), styles["TdCenter"]),
                Paragraph(_escape_xml(op["service"]), styles["Td"]),
                Paragraph(_escape_xml(op["client"] or "-"), styles["Td"]),
                Paragraph(_escape_xml(op["identifier"] or "-"), styles["Td"]),
                Paragraph(_fmt_money(amt), styles["TdRight"]),
                Paragraph(op["currency"], styles["TdCenter"]),
                Paragraph(_fmt_money(fee), styles["TdRight"]),
                Paragraph(_escape_xml(op["agent"]), styles["Td"]),
            ])

        op_data.append([
            Paragraph("TOTAL CONSOLIDÉ DES OPÉRATIONS", styles["TotalLabel"]),
            Paragraph(f"{len(operations)} ops", styles["TotalCenter"]),
            Paragraph("", styles["TotalCenter"]),
            Paragraph("", styles["TotalLabel"]),
            Paragraph("", styles["TotalLabel"]),
            Paragraph("", styles["TotalLabel"]),
            Paragraph(_fmt_money(tot_op_amt), styles["TotalValue"]),
            Paragraph(op_curr, styles["TotalCenter"]),
            Paragraph(_fmt_money(tot_op_fees), styles["TotalValue"]),
            Paragraph("Commissions validées", styles["TotalLabel"]),
        ])

        op_widths = [66, 56, 50, 62, 115, 85, 68, 32, 58, usable_width - (66 + 56 + 50 + 62 + 115 + 85 + 68 + 32 + 58)]
        op_table = Table(op_data, colWidths=op_widths, repeatRows=1, hAlign="LEFT")
        op_table.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), c_primary),
            ("GRID", (0, 0), (-1, -2), 0.5, c_border),
            ("ROWBACKGROUNDS", (0, 1), (-1, -2), [c_white, c_row_alt]),
            ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
            ("TOPPADDING", (0, 0), (-1, -1), 2),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 2),
            ("LEFTPADDING", (0, 0), (-1, -1), 3),
            ("RIGHTPADDING", (0, 0), (-1, -1), 3),
            ("BACKGROUND", (0, -1), (-1, -1), c_dark),
            ("LINEABOVE", (0, -1), (-1, -1), 1, c_gold),
            ("SPAN", (0, -1), (0, -1)),
        ]))
        story.append(op_table)

        # Annexe B: Table de correspondance des références uniques d'opérations (UUID)
        story.append(PageBreak())
        story.append(Paragraph("Annexe B : Table de correspondance des références uniques d'opérations (UUID)", styles["SectionHeading"]))
        annex_b_headers = [
            Paragraph("Repère", styles["ThCenter"]),
            Paragraph("Date / Heure", styles["ThCenter"]),
            Paragraph("Client / Bénéficiaire", styles["Th"]),
            Paragraph("Montant", styles["ThRight"]),
            Paragraph("Dev.", styles["ThCenter"]),
            Paragraph("Référence unique complète (UUID système)", styles["Th"]),
        ]
        annex_b_data = [annex_b_headers]
        for op in operations:
            ref_raw = op["reference"]
            ref_display = f"[{ref_raw[:8]}]" if len(ref_raw) > 12 else ref_raw
            annex_b_data.append([
                Paragraph(ref_display, styles["TdCenter"]),
                Paragraph(op["date"], styles["TdCenter"]),
                Paragraph(_escape_xml(op["client"] or "-"), styles["Td"]),
                Paragraph(_fmt_money(op["amount"]), styles["TdRight"]),
                Paragraph(op["currency"], styles["TdCenter"]),
                Paragraph(f"<font name='Courier' size=6.5>{ref_raw}</font>", styles["Td"]),
            ])
        b_widths = [55, 75, 140, 65, 32, usable_width - (55 + 75 + 140 + 65 + 32)]
        annex_b_table = Table(annex_b_data, colWidths=b_widths, repeatRows=1, hAlign="LEFT")
        annex_b_table.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), c_primary),
            ("GRID", (0, 0), (-1, -1), 0.5, c_border),
            ("ROWBACKGROUNDS", (0, 1), (-1, -1), [c_white, c_row_alt]),
            ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
            ("TOPPADDING", (0, 0), (-1, -1), 2),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 2),
            ("LEFTPADDING", (0, 0), (-1, -1), 3),
            ("RIGHTPADDING", (0, 0), (-1, -1), 3),
        ]))
        story.append(annex_b_table)
    else:
        empty_box = Table(
            [[Paragraph("<i>Aucune opération enregistrée sur cette période.</i>", styles["NoticeEmpty"])]],
            colWidths=[usable_width], hAlign="LEFT"
        )
        empty_box.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, -1), c_light),
            ("BOX", (0, 0), (-1, -1), 0.5, c_border),
            ("PADDING", (0, 0), (-1, -1), 4),
        ]))
        story.append(empty_box)

    canvas_factory = lambda *args, **kwargs: _MonthlyNumberedCanvas(
        *args, downloader=downloader_name, logo_path=logo_path, **kwargs
    )
    doc.build(story, canvasmaker=canvas_factory)
    return stream.getvalue()


def generate_monthly_report(*, user, year, month, format, agent=None, stakeholder=None, stakeholder_type=""):
    snapshot = monthly_financial_snapshot(
        user=user, year=year, month=month, agent=agent,
        stakeholder=stakeholder, stakeholder_type=stakeholder_type,
    )
    format = format.upper()
    builders = {"CSV": _monthly_csv, "XLSX": _monthly_xlsx, "PDF": _monthly_pdf}
    if format not in builders: raise ValueError("Format non pris en charge.")
    export = ReportExport.objects.create(requested_by=user, kind="MONTHLY_FINANCIAL", format=format, filters=snapshot["filters"] | snapshot["period"], snapshot=snapshot)
    try:
        content = _monthly_pdf(snapshot, user) if format == "PDF" else builders[format](snapshot)
        export.file.save(f"rapport-mensuel-{year}-{int(month):02d}.{format.lower()}", ContentFile(content), save=False)
        export.status = "READY"; export.completed_at = timezone.now(); export.save(update_fields=["file", "status", "completed_at"])
    except Exception:
        if export.file:
            export.file.delete(save=False)
        export.status = "FAILED"; export.completed_at = timezone.now(); export.save(update_fields=["status", "completed_at"])
        raise
    record(actor=user, action="MONTHLY_REPORT_GENERATE", instance=export, after={"period": snapshot["period"], "filters": snapshot["filters"]})
    return export


# --------------------------------------------------------------------------- finance report centre

def _text(value):
    if isinstance(value, Decimal):
        return f"{value.quantize(Decimal('0.01')):,.2f}".replace(",", " ")
    return value.strftime("%d/%m/%Y") if hasattr(value, "strftime") else ("" if value is None else str(value))


def _finance_csv(report):
    from apps.finance.reporting import header_rows
    stream = StringIO(newline="")
    writer = csv.writer(stream)
    for row in header_rows(report):
        writer.writerow(_safe_row(row))
    for section in report["sections"]:
        writer.writerow([])
        writer.writerow([section["title"].upper()])
        writer.writerow(section["columns"])
        for row in section["rows"]:
            writer.writerow(_safe_row([_text(v) if not isinstance(v, Decimal) else v for v in row]))
        for row in section["totals"]:
            writer.writerow(_safe_row([_text(v) if not isinstance(v, Decimal) else v for v in row]))
    return ("\ufeff" + stream.getvalue()).encode("utf-8")


def _finance_xlsx(report):
    from apps.finance.reporting import header_rows
    workbook = Workbook()
    summary = workbook.active
    summary.title = "En-tête"
    summary.append(["Champ", "Valeur"])
    for row in header_rows(report):
        summary.append(_safe_row(row))
    for index, section in enumerate(report["sections"], start=1):
        sheet = workbook.create_sheet(f"{index}. {section['title']}"[:31].replace("/", "-").replace(":", ""))
        sheet.append(section["columns"])
        for row in section["rows"]:
            sheet.append(_safe_row([v if isinstance(v, (Decimal, int)) else _text(v) for v in row]))
        for row in section["totals"]:
            sheet.append(_safe_row([v if isinstance(v, (Decimal, int)) else _text(v) for v in row]))
            for cell in sheet[sheet.max_row]:
                cell.font = Font(bold=True)
        for column_index in section["numeric"]:
            for cell in list(sheet.columns)[column_index][1:]:
                cell.number_format = "#,##0.00"
                cell.alignment = Alignment(horizontal="right")
    _style_workbook(workbook)
    stream = BytesIO()
    workbook.save(stream)
    return stream.getvalue()


def _finance_pdf(report, user):
    from apps.finance.reporting import header_rows
    sections = [("En-tête et synthèse", [["Champ", "Valeur"], *header_rows(report)])]
    for section in report["sections"]:
        data = [section["columns"], *[[_text(v) for v in row] for row in section["rows"]]]
        data.extend([[_text(v) for v in row] for row in section["totals"]])
        sections.append((section["title"], data, section["numeric"]))
    subtitle = f"{report['period']['label']} · {report['period']['timezone']} · {report['status']} · {report['id']} v{report['version']}"
    return _build_branded_pdf(title=report["title"], subtitle=subtitle, sections=sections, downloaded_by=user, landscape_mode=True)


def generate_finance_report(*, user, kind, preset, anchor, format, custom_end=None, filters=None):
    from apps.finance.reporting import build_report, can_download
    from django.core.exceptions import PermissionDenied
    if not can_download(user):
        raise PermissionDenied("Le téléchargement des rapports n’est pas activé pour votre rôle.")
    report = build_report(user=user, kind=kind, preset=preset, anchor=anchor, custom_end=custom_end, filters=filters)
    format = format.upper()
    builders = {"CSV": _finance_csv, "XLSX": _finance_xlsx}
    if format not in {*builders, "PDF"}:
        raise ValueError("Format non pris en charge.")
    export = ReportExport.objects.create(requested_by=user, kind=f"FINANCE_{kind}", format=format,
        filters={"start_date": str(report["period"]["start"]), "end_date": str(report["period"]["end"]), "preset": preset, **report["filters"]},
        snapshot={"id": report["id"], "version": report["version"], "status": report["status"],
                  "sections": [{"title": s["title"], "columns": s["columns"], "rows": [[_text(v) for v in r] for r in s["rows"]], "totals": [[_text(v) for v in r] for r in s["totals"]]} for s in report["sections"]]})
    try:
        content = _finance_pdf(report, user) if format == "PDF" else builders[format](report)
        export.file.save(f"{kind.lower()}-{report['period']['start']}-{report['period']['end']}.{format.lower()}", ContentFile(content), save=False)
        export.status, export.completed_at = "READY", timezone.now()
        export.save(update_fields=["file", "status", "completed_at"])
    except Exception:
        if export.file:
            export.file.delete(save=False)
        export.status, export.completed_at = "FAILED", timezone.now()
        export.save(update_fields=["status", "completed_at"])
        raise
    record(actor=user, action="REPORT_GENERATE", instance=export, after={"kind": kind, "format": format, "report_id": report["id"]})
    return export
