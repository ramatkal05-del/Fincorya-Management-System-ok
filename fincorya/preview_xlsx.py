"""Generate XLSX previews and verify totals match dashboard snapshots."""
import os
from datetime import date
from unittest.mock import patch
from zoneinfo import ZoneInfo

os.environ.setdefault("DJANGO_SETTINGS_MODULE", "config.settings")
import django
django.setup()

from openpyxl import load_workbook

from apps.accounts.models import Role, User
from apps.finance.reporting import build_report
from apps.reports.services import _xlsx_bytes, operation_report_snapshot, monthly_financial_snapshot

OUT = "output/xlsx"
os.makedirs(OUT, exist_ok=True)

admin = User.objects.filter(role=Role.ADMIN, is_active=True).order_by("id").first()
manager = User.objects.filter(role=Role.FINANCE_MANAGER, is_active=True).order_by("id").first()
agent = User.objects.filter(role=Role.AGENT, is_active=True).order_by("id").first()

Z = ZoneInfo("Europe/Istanbul")
when = __import__("datetime").datetime(2026, 9, 15, 16, tzinfo=Z)

with patch("django.utils.timezone.now", return_value=when):
    # Agent daily XLSX
    agent_snap = operation_report_snapshot(user=agent, start_date=date(2026, 9, 15), end_date=date(2026, 9, 15))
    agent_xlsx = _xlsx_bytes(agent_snap)
    open(f"{OUT}/preview-agent-daily.xlsx", "wb").write(agent_xlsx)
    print("Agent XLSX: rows=", agent_snap["totals"]["count"], "by_currency=", agent_snap["totals"]["by_currency"])

    # Finance weekly XLSX
    week_snap = build_report(user=manager, kind="TREASURY", preset="WEEK", anchor=date(2026, 9, 15))
    # build_report returns a different structure - use operation snapshot for XLSX
    week_op = operation_report_snapshot(user=manager, start_date=date(2026, 9, 14), end_date=date(2026, 9, 20))
    week_xlsx = _xlsx_bytes(week_op)
    open(f"{OUT}/preview-finance-weekly.xlsx", "wb").write(week_xlsx)
    print("Finance weekly XLSX: rows=", week_op["totals"]["count"], "by_currency=", week_op["totals"]["by_currency"])

    # Admin monthly XLSX
    month_op = operation_report_snapshot(user=admin, start_date=date(2026, 9, 1), end_date=date(2026, 9, 30))
    month_xlsx = _xlsx_bytes(month_op)
    open(f"{OUT}/preview-admin-monthly.xlsx", "wb").write(month_xlsx)
    print("Admin monthly XLSX: rows=", month_op["totals"]["count"], "by_currency=", month_op["totals"]["by_currency"])

# Verify XLSX content by reading back
for name in ("preview-agent-daily.xlsx", "preview-finance-weekly.xlsx", "preview-admin-monthly.xlsx"):
    path = os.path.join(OUT, name)
    wb = load_workbook(path)
    print(f"\n=== {name} ===")
    print(f"Sheets: {wb.sheetnames}")
    detail = wb["Détail"]
    print(f"Détail rows: {detail.max_row}, cols: {detail.max_column}")
    print(f"Freeze panes: {detail.freeze_panes}")
    print(f"Gridlines: {detail.sheet_view.showGridLines}")
    print(f"Auto filter: {detail.auto_filter.ref}")
    # Header style
    h = detail["A1"]
    print(f"Header fill: {h.fill.fgColor.rgb}, font color: {h.font.color.rgb if h.font.color else None}, bold: {h.font.bold}")
    # Synth sheet
    synth = wb["Synthèse"]
    print(f"Synthèse rows: {synth.max_row}")
    for row in synth.iter_rows(min_row=1, max_row=min(synth.max_row, 6), values_only=True):
        print(f"  {row}")
    # Properties
    print(f"Creator: {wb.properties.creator}, Title: {wb.properties.title}")
    wb.close()

# Verify totals match
print("\n=== Totals comparison ===")
agent_total = sum(float(v["amount"]) for v in agent_snap["totals"]["by_currency"].values())
print(f"Agent snapshot total: {agent_total}")
month_total = sum(float(v["amount"]) for v in month_op["totals"]["by_currency"].values())
print(f"Month snapshot total: {month_total}")

# Read back XLSX totals
wb_agent = load_workbook(os.path.join(OUT, "preview-agent-daily.xlsx"))
synth_agent = wb_agent["Synthèse"]
xlsx_agent_total = 0
for row in synth_agent.iter_rows(min_row=2, values_only=True):
    if row[0] and "Montant total" in str(row[0]):
        xlsx_agent_total += float(row[1])
print(f"Agent XLSX total: {xlsx_agent_total}")
print(f"Match: {abs(agent_total - xlsx_agent_total) < 0.01}")
wb_agent.close()
